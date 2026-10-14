import io
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def clean_markdown(text: str) -> str:
    """Removes markdown symbols for plain-text formatting."""
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"#{1,6}\s?", "", text)
    return text.strip()


def extract_paragraphs(text: str, max_items: int = 4) -> list:
    """Splits markdown notes into substantive bullet points."""
    cleaned = clean_markdown(text)
    items = []
    for line in cleaned.split("\n"):
        line = line.strip()
        if len(line) > 25 and not line.startswith("http"):
            items.append(line)
    return items[:max_items] if items else [cleaned[:250]]


# -------------------------------------------------------------
# 1. WORD DOCUMENT GENERATOR (All details shown on site)
# -------------------------------------------------------------
def generate_word_document(topic: str, report: str, research_notes: str, analysis: str, sources: list) -> bytes:
    doc = Document()

    title = doc.add_heading("Comprehensive Research Briefing", level=0)
    title.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    p_topic = doc.add_paragraph()
    p_topic.add_run("Research Topic: ").bold = True
    p_topic.add_run(topic)
    doc.add_paragraph()

    doc.add_heading("1. Executive Research Report", level=1)
    doc.add_paragraph(clean_markdown(report))
    doc.add_paragraph()

    doc.add_heading("2. Research Analyst's Insights", level=1)
    doc.add_paragraph(clean_markdown(analysis))
    doc.add_paragraph()

    doc.add_heading("3. Raw Research Findings & Evidence", level=1)
    doc.add_paragraph(clean_markdown(research_notes))
    doc.add_paragraph()

    doc.add_heading("4. Verified Sources", level=1)
    for item in sources:
        title_text = item.get("title", "Source")
        link_text = item.get("link", "")
        doc.add_paragraph(f"[{item.get('id', '-')}] {title_text} — {link_text}", style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------
# 2. EXECUTIVE CONSULTING-GRADE POWERPOINT (.pptx)
# -------------------------------------------------------------
def generate_pptx_deck(topic: str, report: str, analysis: str, sources: list) -> bytes:
    prs = Presentation()
    # 16:9 Modern Widescreen
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette
    NAVY = PptxRGBColor(15, 23, 42)        # #0F172A
    CARD_BG = PptxRGBColor(241, 245, 249)  # #F1F5F9
    ACCENT_BLUE = PptxRGBColor(37, 99, 235)# #2563EB
    TEXT_DARK = PptxRGBColor(30, 41, 59)   # #1E293B
    TEXT_MUTED = PptxRGBColor(100, 116, 139)# #64748B
    WHITE = PptxRGBColor(255, 255, 255)
    LIGHT_BORDER = PptxRGBColor(203, 213, 225)

    # ---------------------------------------------------------
    # SLIDE 1: Premium Dark Hero Cover Slide
    # ---------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptxInches(13.333), PptxInches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Left Accent Bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptxInches(1.0), PptxInches(1.8), PptxInches(0.12), PptxInches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_BLUE
    bar.line.fill.background()

    # Title & Subtitle Frame
    t_box = s1.shapes.add_textbox(PptxInches(1.4), PptxInches(1.6), PptxInches(10.5), PptxInches(4.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    tag_p = tf1.paragraphs[0]
    tag_p.text = "EXECUTIVE STRATEGIC BRIEFING"
    tag_p.font.size = PptxPt(12)
    tag_p.font.bold = True
    tag_p.font.color.rgb = ACCENT_BLUE

    title_p = tf1.add_paragraph()
    title_p.text = topic.strip()
    title_p.font.size = PptxPt(36)
    title_p.font.bold = True
    title_p.font.color.rgb = WHITE
    title_p.space_before = PptxPt(12)

    sub_p = tf1.add_paragraph()
    sub_p.text = "Autonomous Intelligence & Strategic Synthesis Report | Multi-Agent Analysis"
    sub_p.font.size = PptxPt(15)
    sub_p.font.color.rgb = PptxRGBColor(148, 163, 184)
    sub_p.space_before = PptxPt(16)

    # ---------------------------------------------------------
    # Helper: Consistent Light Slide Header
    # ---------------------------------------------------------
    def apply_slide_header(slide, title_text: str, subtitle_text: str):
        box = slide.shapes.add_textbox(PptxInches(1.0), PptxInches(0.55), PptxInches(11.33), PptxInches(1.2))
        tf = box.text_frame
        p_sub = tf.paragraphs[0]
        p_sub.text = subtitle_text.upper()
        p_sub.font.size = PptxPt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = ACCENT_BLUE

        p_main = tf.add_paragraph()
        p_main.text = title_text
        p_main.font.size = PptxPt(26)
        p_main.font.bold = True
        p_main.font.color.rgb = NAVY
        p_main.space_before = PptxPt(4)

    # ---------------------------------------------------------
    # SLIDE 2: 3-Card Executive Summary Layout
    # ---------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    apply_slide_header(s2, "Executive Findings & Core Takeaways", "Key Synthesis")

    bullets_report = extract_paragraphs(report, 3)
    card_width = PptxInches(3.55)
    card_height = PptxInches(4.8)
    card_gap = PptxInches(0.34)
    start_left = PptxInches(1.0)
    card_top = PptxInches(1.9)

    card_headers = ["Core Proposition", "Market Catalyst", "Strategic Outlook"]

    for i in range(3):
        c_left = start_left + i * (card_width + card_gap)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, card_top, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = LIGHT_BORDER
        card.line.width = PptxPt(1)

        t_box = s2.shapes.add_textbox(c_left + PptxInches(0.2), card_top + PptxInches(0.2), card_width - PptxInches(0.4), card_height - PptxInches(0.4))
        tf = t_box.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = f"0{i+1}"
        p_num.font.size = PptxPt(18)
        p_num.font.bold = True
        p_num.font.color.rgb = ACCENT_BLUE

        p_head = tf.add_paragraph()
        p_head.text = card_headers[i]
        p_head.font.size = PptxPt(16)
        p_head.font.bold = True
        p_head.font.color.rgb = NAVY
        p_head.space_after = PptxPt(10)

        p_body = tf.add_paragraph()
        p_body.text = bullets_report[i] if i < len(bullets_report) else "Synthesized strategic data point available in full report."
        p_body.font.size = PptxPt(12)
        p_body.font.color.rgb = TEXT_DARK
        p_body.space_before = PptxPt(6)

    # ---------------------------------------------------------
    # SLIDE 3: 2-Column Split (Drivers vs. Risks)
    # ---------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    apply_slide_header(s3, "Analyst Evaluation & Strategic Friction", "Deep Dive Assessment")

    col_width = PptxInches(5.45)
    col_height = PptxInches(4.8)
    col_gap = PptxInches(0.43)
    col_top = PptxInches(1.9)

    analysis_points = extract_paragraphs(analysis, 4)
    half = max(1, len(analysis_points) // 2)
    left_points = analysis_points[:half]
    right_points = analysis_points[half:]

    # Left Column: Growth Catalysts & Trends
    c_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(1.0), col_top, col_width, col_height)
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = CARD_BG
    c_left.line.color.rgb = LIGHT_BORDER
    c_left.line.width = PptxPt(1)

    tf_l = s3.shapes.add_textbox(PptxInches(1.2), col_top + PptxInches(0.2), col_width - PptxInches(0.4), col_height - PptxInches(0.4)).text_frame
    tf_l.word_wrap = True
    p1 = tf_l.paragraphs[0]
    p1.text = "📈 Growth Drivers & Industry Trends"
    p1.font.size = PptxPt(16)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    p1.space_after = PptxPt(12)

    for pt in left_points:
        p = tf_l.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = PptxPt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_after = PptxPt(8)

    # Right Column: Risks & Uncertainties
    c_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(1.0) + col_width + col_gap, col_top, col_width, col_height)
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = CARD_BG
    c_right.line.color.rgb = LIGHT_BORDER
    c_right.line.width = PptxPt(1)

    tf_r = s3.shapes.add_textbox(PptxInches(1.2) + col_width + col_gap, col_top + PptxInches(0.2), col_width - PptxInches(0.4), col_height - PptxInches(0.4)).text_frame
    tf_r.word_wrap = True
    p2 = tf_r.paragraphs[0]
    p2.text = "⚠️ Friction Points, Risks & Countermeasures"
    p2.font.size = PptxPt(16)
    p2.font.bold = True
    p2.font.color.rgb = NAVY
    p2.space_after = PptxPt(12)

    for pt in right_points:
        p = tf_r.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = PptxPt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_after = PptxPt(8)

    # ---------------------------------------------------------
    # SLIDE 4: Primary Evidence & Verified Sources Grid
    # ---------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    apply_slide_header(s4, "Primary Corpus & Research Citations", "Verification Matrix")

    top_pos = PptxInches(1.9)
    for i, s in enumerate(sources[:4]):
        row_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptxInches(1.0), top_pos, PptxInches(11.33), PptxInches(1.05))
        row_card.fill.solid()
        row_card.fill.fore_color.rgb = CARD_BG
        row_card.line.color.rgb = LIGHT_BORDER
        row_card.line.width = PptxPt(1)

        tf_s = s4.shapes.add_textbox(PptxInches(1.2), top_pos + PptxInches(0.12), PptxInches(10.9), PptxInches(0.8)).text_frame
        tf_s.word_wrap = True

        p_t = tf_s.paragraphs[0]
        p_t.text = f"[{s.get('id', i+1)}] {s.get('title', 'Verified Web Citation')}"
        p_t.font.size = PptxPt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

        p_u = tf_s.add_paragraph()
        p_u.text = s.get("link", "Direct web archive")
        p_u.font.size = PptxPt(10)
        p_u.font.color.rgb = ACCENT_BLUE

        top_pos += PptxInches(1.2)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------
# 3. EXECUTIVE BRIEFING PDF GENERATOR (.pdf)
# -------------------------------------------------------------
def generate_pdf_report(topic: str, report: str, analysis: str) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(0, 10, "Executive Research Presentation Report", ln=True)

    pdf.set_font("Helvetica", "I", 11)
    pdf.set_text_color(100, 116, 139)
    topic_sanitized = topic.encode("latin-1", "replace").decode("latin-1")
    pdf.cell(0, 8, f"Topic: {topic_sanitized}", ln=True)
    pdf.ln(4)

    pdf.set_draw_color(203, 213, 225)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(6)

    # Section 1: Executive Briefing
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(0, 8, "1. Executive Briefing", ln=True)
    pdf.ln(2)

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(30, 41, 59)
    clean_rep = clean_markdown(report).encode("latin-1", "replace").decode("latin-1")
    for line in clean_rep.split("\n"):
        if line.strip():
            pdf.multi_cell(0, 5, line.strip())
            pdf.ln(2)

    pdf.ln(4)

    # Section 2: Strategic Analysis
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(0, 8, "2. Strategic Analyst Evaluation", ln=True)
    pdf.ln(2)

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(30, 41, 59)
    clean_an = clean_markdown(analysis).encode("latin-1", "replace").decode("latin-1")
    for line in clean_an.split("\n"):
        if line.strip():
            pdf.multi_cell(0, 5, line.strip())
            pdf.ln(2)

    return bytes(pdf.output())